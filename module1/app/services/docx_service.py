from io import BytesIO


def extract_docx_text(content: bytes) -> str:
    from docx import Document
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(paragraphs).strip()
