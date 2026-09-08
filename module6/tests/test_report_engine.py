"""End-to-end: fixture entities across Modules 1-5 -> a complete generated
report (.docx + manifest), via the same ReportEngine the API route uses.
Covers both a complete fixture set and a deliberately incomplete one, per
the spec's explicit dual-case testing requirement."""

from pathlib import Path

from app.schemas.report import ReportStatus


async def test_complete_fixture_set_generates_a_full_report(complete_report_engine, tmp_path):
    report = await complete_report_engine.generate_report("vendor:report-complete", version=1)

    assert report is not None
    assert report.status == ReportStatus.DRAFT
    assert report.version == 1
    assert report.missing_sections == []
    assert report.label.startswith("AI-drafted")

    generated_file = tmp_path / "reports" / report.file_path
    assert generated_file.exists()
    assert generated_file.suffix == ".docx"

    from docx import Document

    doc = Document(str(generated_file))
    assert len(doc.tables) == 3  # screening, adverse media, triage
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Meridian Capital Holdings Sdn Bhd" in all_text
    assert "PLACEHOLDER TEMPLATE" in all_text


async def test_incomplete_fixture_set_still_generates_a_report_with_flagged_sections(incomplete_report_engine, tmp_path):
    report = await incomplete_report_engine.generate_report("vendor:report-incomplete", version=1)

    assert report is not None
    assert set(report.missing_sections) == {
        "screening_findings_summary", "adverse_media_findings", "triage_summary", "risk_assessment",
    }

    generated_file = tmp_path / "reports" / report.file_path
    assert generated_file.exists()

    from docx import Document

    doc = Document(str(generated_file))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "No Module 2 screening evidence found" in all_text
    assert "No Module 3 adverse media findings found" in all_text
    assert "No Module 5 risk assessment found" in all_text


async def test_unknown_entity_returns_none(complete_report_engine):
    report = await complete_report_engine.generate_report("vendor:does-not-exist", version=1)
    assert report is None


async def test_manifest_and_report_share_the_same_report_id(complete_report_engine):
    report = await complete_report_engine.generate_report("vendor:report-complete", version=1)
    assert report.manifest.report_id == report.report_id
    assert report.manifest.entity_id == report.entity_id
    assert report.manifest.version == report.version
